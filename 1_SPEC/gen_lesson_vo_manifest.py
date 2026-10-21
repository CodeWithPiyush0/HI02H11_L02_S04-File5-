# -*- coding: utf-8 -*-
"""VO manifest for HI02H11_L02_S04 (ओ / औ), in the order the child hears it - for the SME to record.

Ported from File3's (HI02H11_L02_S02) 1_SPEC/gen_lesson_vo_manifest.py: the same sheets and columns.

ORDER. Start screen, transition 1, pages 1-5, transition 2, pages 6-16, transition 3, page 17 (the
«मात्रा टोकरी» game), page 18 (celebration). Inside a page the clips are in the order they PLAY, taken
from a real play-through of the built lesson (1_SPEC/vo_play_order/lesson_play_order.json, recorded by
the two drive_*.py scripts beside it): what plays as the page opens, then for each item its 1st / 2nd /
3rd wrong try (Hint 1, 2, 3) and its correct answer. A clip that plays again later on the same page is
listed once, at its first play. Clips a child reaches only with other choices than the play-through
made are listed at the end of their page.

TEXT. Every line is the build's own card.json text (assets.audio_text); the transitions are the shared
kit recordings, whose lines are given below.

Run:   PYTHONUTF8=1 python 1_SPEC/gen_lesson_vo_manifest.py
Re-record the play order after any change to a page's flow (serve 3_BUILD on :8979):
       python 1_SPEC/vo_play_order/drive_ladder_pass.py    http://localhost:8979/HI02H11_L02_S04.html ladder.json
       python 1_SPEC/vo_play_order/drive_first_try_pass.py http://localhost:8979/HI02H11_L02_S04.html first.json
       then merge them as {"ladder_pass": ..., "first_try_pass": ...} into lesson_play_order.json
"""
import io, os, json, subprocess, collections
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(HERE, "3_BUILD")
OUT = os.path.join(HERE, "HI02H11_L02_S04_VO_manifest.xlsx")      # the folder root, for the SME
PO = json.load(io.open(os.path.join(HERE, "1_SPEC", "vo_play_order", "lesson_play_order.json"), encoding="utf-8"))
LAD, FIRST = PO["ladder_pass"], PO["first_try_pass"]

card = json.load(io.open(os.path.join(BUILD, "card.json"), encoding="utf-8"))
slides = card["slides"]
declared = card["assets"]["audio"]
text = dict(card["assets"].get("audio_text") or {})
KIT_TEXT = {"vo_pt_tutorial": "ध्यान से देखिए और मेरे साथ जानिए। चलिए, शुरू करें!",
            "vo_pt_guided": "बहुत बढ़िया! अब हम साथ मिलकर शुरू करते हैं। चलिए, साथ में करें!",
            "vo_pt_practice": "वाह! अब आपकी बारी।"}
for k, v in KIT_TEXT.items():
    text.setdefault(k, v)

# Lines that are already File3's RECORDED narrator (the same line, copied byte for byte from File3).
# Everything else is a Gemini TTS placeholder (voice Leda) waiting for the SME's recording.
HUMAN = {"vo_pt_tutorial", "vo_pt_guided", "vo_pt_practice", "vo_mic_nudge", "vo_cel_prompt", "vo_g4_h1",
         "vo_g4_prompt", "vo_g4d_end", "vo_g5_h1", "vo_sc_h2", "vo_sc_prompt", "vo_try_again"}

# [2026-10-21] the SME's recordings, as delivered (1_SPEC/lesson_voice/source_wav/<id>.wav)
SME = {os.path.splitext(f)[0] for f in os.listdir(os.path.join(HERE, "1_SPEC", "lesson_voice", "source_wav"))
       if f.lower().endswith(".wav")} if os.path.isdir(os.path.join(HERE, "1_SPEC", "lesson_voice", "source_wav")) else set()

SCREEN = {"MATRA_PAIRS": "Meet the pair", "MATRA_BUILD": "Build the word", "MEET_PAIR": "Word examples",
          "TRAIN_TAP": "Tap the coach", "TRAIN_SORT": "Sort into coaches", "WORD_BUILD": "Complete the word",
          "SENTENCE_COMPLETE": "Complete the sentence", "MATRA_TOKRI": "«मात्रा टोकरी» game",
          "CELEBRATION": "Celebration"}

# ---- what each clip IS on a page, from the card ------------------------------------------------
ROLE = {"prompt": "page instruction", "hint1": "Hint 1", "hint2": "Hint 2", "hint3": "Hint 3 (names the answer)",
        "correct": "correct-answer praise", "outro": "closing line", "base": "the base word",
        "onset": "the letter takes the matra", "matra_name": "names the matra", "result": "the new word",
        "sounds": "sounds the word out"}
FIELD = {"pairs.audio": "the letter and its matra", "examples.audio_line": "example word, with its matra",
         "coaches.audio": "coach word", "bins.audio": "coach label", "cards.audio": "card name",
         "cards.correct_audio": "correct - this card", "cards.hint2_audio": "Hint 2 - this card",
         "cards.hint3_audio": "Hint 3 - this card", "options.audio": "option",
         "slots.name_audio": "picture name", "slots.correct_audio": "correct - this word",
         "slots.hint3_audio": "Hint 3 - this word", "options.sentence_audio": "the sentence read with this option",
         "rounds.vo": "round instruction", "rounds.cheer": "round praise + next round's instruction"}


def what_is(page, cid):
    s = slides[page - 1]
    for k, v in (s.get("audio") or {}).items():
        if v == cid:
            return ROLE.get(k, k)
    d = s.get("data") or {}
    if d.get("mic_nudge_audio") == cid:
        return "mic reminder (after 7 s with no tap)"
    if d.get("done_audio") == cid:
        return "game complete"
    if cid.startswith("vo_mt_w_"):
        return "a word in the basket game"
    if isinstance(d.get("phonemes"), dict) and cid in d["phonemes"].values():
        return "the matra's sound"
    for grp, items in d.items():
        if not isinstance(items, list):
            continue
        for it in items:
            if not isinstance(it, dict):
                continue
            for f, v in it.items():
                if v == cid:
                    nm = next((it[k] for k in ("word", "letter", "label", "akshar") if isinstance(it.get(k), str)), "")
                    if s["type"] == "WORD_BUILD" and f == "audio":
                        return "letter card «%s»" % nm
                    return FIELD.get("%s.%s" % (grp, f), "%s.%s" % (grp, f)) + (" «%s»" % nm if nm else "")
    if cid.startswith("vo_name_"):
        return "name of a picture / word on this page"
    if cid.startswith("vo_matra_"):
        return "the matra's name"
    if cid.startswith("vo_letter_"):
        return "coach label"
    if cid == "vo_try_again":
        return "try again"
    return ""


def page_ids(p):
    """every clip the page's own card entry names"""
    s = slides[p - 1]
    out = list((s.get("audio") or {}).values())

    def walk(x):
        if isinstance(x, dict):
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
        elif isinstance(x, str) and x.startswith("vo_") and x in declared:
            out.append(x)
    walk(s.get("data") or {})
    return [c for c in dict.fromkeys(out) if c]


# ---- the play order ---------------------------------------------------------------------------
def page_steps(p):
    """[(when, [clips...]), ...] for page p, in play order"""
    s, d = slides[p - 1], slides[p - 1].get("data") or {}
    lp = LAD["pages"][str(p)]
    fp = FIRST.get(str(p))
    entry = list(fp["entry"] if fp else lp["entry"])
    if s["type"] == "TRAIN_SORT" and not d.get("demo"):
        # the tray is shuffled on every visit, so its names are listed in the card's own order
        names = [c.get("audio") for c in d.get("cards", [])]
        got = [c for c in entry if c in names]
        rest = [c for c in entry if c not in names]
        entry = rest[:1] + [n for n in names if n in got] + rest[1:]
    if d.get("demo"):
        when0 = "Watch-only demo: the hand moves each card while it is named and explained"
    elif s["type"] in ("MATRA_PAIRS", "MATRA_BUILD", "MEET_PAIR"):
        when0 = "Page opens - the teaching sequence plays by itself"
    elif s["type"] == "TRAIN_SORT":
        when0 = "Page opens; the cards come into the tray, each one named (tray order is shuffled)"
    else:
        when0 = "Page opens"
    if s["type"] == "MEET_PAIR":
        # the page waits for the child's mic tap after each word, so the play-through heard only the
        # first word; the words are listed in their order from the card
        entry = [c for c in entry if c != d.get("mic_nudge_audio")]
        steps = [(when0, entry)]
        for k, ex in enumerate(d.get("examples", [])[1:], 2):
            steps.append(("Word %d «%s» - after the child taps the mic for word %d" % (k, ex["word"], k - 1),
                          [ex["audio_line"]]))
        if d.get("mic_nudge_audio"):
            steps.append(("After each word: if the mic is not tapped for 7 s (and every 7 s after)",
                          [d["mic_nudge_audio"]]))
        return steps
    steps = [(when0, entry)]
    right = {x["item"]: x["clips"] for x in (fp or {}).get("right", [])}
    for it in lp["items"]:
        nm = it["item"]
        tag = (" - «%s»" % nm) if len(lp["items"]) > 1 else ""
        prompt = (s.get("audio") or {}).get("prompt")
        for k, lab in (("miss1", "1st wrong try"), ("miss2", "2nd wrong try"), ("miss3", "3rd wrong try")):
            # the page's own instruction inside a hint step is the idle reminder replaying it while
            # the play-through waited - not part of the hint
            clips = [c for c in it.get(k) or [] if c != prompt]
            if clips:
                steps.append(("%s%s" % (lab, tag), clips))
        if right.get(nm):
            steps.append(("Correct answer%s" % tag, right[nm]))
        if it.get("right_after_hints"):
            steps.append(("Correct answer after Hint 3%s" % tag, it["right_after_hints"]))
    return steps


seq = [("Start", "Start screen", "Start screen", [("Start screen - Swifty's greeting (and the speaker button replays it)", LAD["start"])])]
gate_at = {}
pin = (card.get("gate") or {}).get("at") or {}
for ph in ("tutorial", "guided", "practice"):
    rnd = {"practice": "round3"}.get(ph, ph)
    at = next((i + 1 for i, s in enumerate(slides) if s["id"] == pin.get(rnd)), None) \
        or next(i + 1 for i, s in enumerate(slides) if s["phase"] == ph)
    gate_at[at] = ph
TITLE = {"tutorial": "«चलिए, शुरू करें!»", "guided": "«चलिए, साथ में करें!»", "practice": "«अब आपकी बारी!»"}
GATE_CLIPS = {"tutorial": [c for c in LAD["gate_tutorial"] if c.startswith("vo_pt_")] or ["vo_pt_tutorial"],
              "guided": LAD.get("gate_guided") or ["vo_pt_guided"],
              "practice": LAD.get("gate_practice") or ["vo_pt_practice"]}
n_gate = 0
for p in range(1, len(slides) + 1):
    if p in gate_at:
        n_gate += 1
        ph = gate_at[p]
        seq.append(("Transition %d" % n_gate, "Transition %s" % TITLE[ph],
                    "Transition %d (before page %d)" % (n_gate, p),
                    [("Swifty rises, then says the line; the text types as she says it", GATE_CLIPS[ph])]))
    s = slides[p - 1]
    nm = SCREEN.get(s["type"], s["type"]) + (" (watch-only demo)" if (s.get("data") or {}).get("demo") else "")
    if s["type"] == "MATRA_TOKRI":
        # the game's clips play as the child steers the basket - listed in the order a full game hears
        # them (measured by 1_SPEC/_drive_tokri.py)
        d_ = s["data"]; r0, r1 = d_["rounds"][0], d_["rounds"][1]
        steps = [("Game opens - the tutorial: the hand drags the basket", [s["audio"]["prompt"]]),
                 ("Round 1 (ओ) starts - the instruction", [r0["vo"]]),
                 ("Round 1 - each ओ word is read as the basket catches it", ["vo_mt_w_" + w["id"] for w in r0["words"]]),
                 ("Round 1 complete - praise + round 2's instruction (one clip)", [r1["cheer"]]),
                 ("Round 2 - each औ word is read as the basket catches it", ["vo_mt_w_" + w["id"] for w in r1["words"]]),
                 ("Both rounds complete", [d_["done_audio"]])]
        if r1.get("vo") in declared:
            # File2's game says round 2's instruction inside the cheer; the line alone is its fallback
            steps.append(("Round 2 instruction alone - only if round 2 starts without the praise clip", [r1["vo"]]))
    elif s["type"] == "CELEBRATION":
        steps = [("Swifty jumps and celebrates first, then says the line (lip-synced)", page_steps(p)[0][1])]
    else:
        steps = page_steps(p)
    # a transition line logged at the end of the page before its gate belongs to the gate row
    steps = [(w, [c for c in cl if not c.startswith("vo_pt_")]) for w, cl in steps]
    heard = {c for _, cl in steps for c in cl}
    rest = [c for c in page_ids(p) if c not in heard]
    # only the question pages branch on the child's choices; the teaching pages and the demo play
    # themselves, so a clip their data names but never plays is not theirs to record
    branches = s["type"] in ("TRAIN_TAP", "TRAIN_SORT", "WORD_BUILD", "SENTENCE_COMPLETE") \
        and not (s.get("data") or {}).get("demo")
    if rest and branches:
        steps.append(("Depends on the child's choices (another card / option answered wrong, or tapped)", rest))
    seq.append((str(p), nm, "Page %d" % p, steps))

# ---- rows ---------------------------------------------------------------------------------------
def duration(path):
    try:
        o = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                            "-of", "default=nw=1:nk=1", path], capture_output=True, text=True).stdout
        return round(float(o.strip()), 2)
    except Exception:
        return None


DUR = {}
pages_of = collections.defaultdict(list)
rows = []
for pkey, screen, label, steps in seq:
    seen = collections.OrderedDict()
    for when, clips in steps:
        for cid in clips:
            if not cid.startswith("vo_"):
                continue
            if cid in seen:
                if when not in seen[cid]["again"] and when != seen[cid]["when"]:
                    seen[cid]["again"].append(when)
                continue
            seen[cid] = {"when": when, "again": []}
    for k, (cid, v) in enumerate(seen.items(), 1):
        if cid not in DUR:
            path = os.path.join(BUILD, declared.get(cid, "assets/Audio/%s.ogg" % cid))
            DUR[cid] = duration(path) if os.path.exists(path) else None
        pages_of[cid].append(pkey)
        if pkey.isdigit():
            what = what_is(int(pkey), cid)
        elif pkey == "Start":
            what = "greeting"
        else:
            what = "transition line"
        if cid == "vo_cel_prompt":
            what = "celebration line"
        rows.append({"page": pkey, "screen": screen, "label": label, "n": k, "id": cid,
                     "file": os.path.basename(declared.get(cid, cid + ".ogg")), "text": text.get(cid, ""),
                     "what": what, "when": v["when"] + (("  ·  plays again: " + "; ".join(v["again"])) if v["again"] else "")})
for r in rows:
    others = [p for p in pages_of[r["id"]] if p != r["page"]]
    r["also"] = ", ".join(("p" + p) if p.isdigit() else p for p in dict.fromkeys(others))
    d_ = DUR.get(r["id"])
    r["dur"] = d_
    r["status"] = ("MISSING - record" if d_ is None else
                   "Recorded (SME) - in the build" if r["id"] in SME else
                   "Recorded voice (File3) - keep or re-record" if r["id"] in HUMAN else
                   "TTS placeholder - record")

used = set(r["id"] for r in rows)
not_played = sorted(c for c in declared if c.startswith("vo_") and c not in used)

# ---- workbook -----------------------------------------------------------------------------------
thin = Side(style="thin", color="C9CED6"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill("solid", fgColor="1F4E79"); DEV = "Nirmala UI"
BAND = ["FFFFFF", "F3F7FB"]


def sheet(ws, title, sub, head, widths, data, dev_cols=(), wrap_cols=(), center=(), fill=None):
    ws.append([title]); ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.append([sub]); ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws.append([]); ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hfill; cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for vals, meta in data:
        ws.append(vals); r = ws.max_row
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=r, column=c); cell.border = box
            if fill: cell.fill = PatternFill("solid", fgColor=fill(meta))
            cell.alignment = Alignment(vertical="center", wrap_text=(c in wrap_cols),
                                       horizontal="center" if c in center else "left")
            if c in dev_cols:
                cell.font = Font(name=DEV, size=12 if c == dev_cols[0] else 10, bold=(c == dev_cols[0]))
        if meta and meta.get("red"):
            ws.cell(row=r, column=len(head)).font = Font(bold=True, color="B03A2E")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(len(head)), ws.max_row)


band_of, b = {}, 0
for r in rows:
    if r["label"] not in band_of:
        band_of[r["label"]] = BAND[b % 2]; b += 1
GATE_FILL = "EEE8FB"

wb = Workbook()
ws = wb.active; ws.title = "VO Manifest"
uniq = list(dict.fromkeys(r["id"] for r in rows))
n_tts = len([c for c in uniq if c not in HUMAN and c not in SME])
sheet(ws, "मात्राओं की रेल (HI02H11_L02_S04 · ओ / औ) - voice-over manifest, in play order",
      "Start screen -> transitions and pages 1-18 in game order; on each page the clips in the order they play. "
      "%d rows, %d distinct recordings (%d to record). Folder 3_BUILD/assets/Audio/  ·  file name = Audio ID"
      % (len(rows), len(uniq), n_tts),
      ["S.No", "Page", "Screen", "#", "Audio ID", "Hindi script (spoken exactly)", "What it is",
       "When it plays", "Same recording also on", "Duration (s)", "Status"],
      [6, 13, 24, 4, 22, 46, 30, 60, 18, 10, 26],
      [([i, r["label"], r["screen"], r["n"], r["id"], r["text"], r["what"], r["when"], r["also"], r["dur"], r["status"]],
        {"fill": GATE_FILL if r["label"].startswith(("Transition", "Start")) else band_of[r["label"]],
         "red": not r["status"].startswith("Recorded")})
       for i, r in enumerate(rows, 1)],
      dev_cols=(6, 7, 8), wrap_cols=(3, 6, 7, 8, 9), center=(1, 4, 10), fill=lambda m: m["fill"])

# one row per recording, in the order it is first heard
u_rows = []
for k, cid in enumerate(uniq, 1):
    r0 = next(r for r in rows if r["id"] == cid)
    pages = list(dict.fromkeys(r["label"] for r in rows if r["id"] == cid))
    u_rows.append(([k, cid, r0["file"], r0["text"], ", ".join(pages), r0["dur"], r0["status"]],
                   {"red": not r0["status"].startswith("Recorded")}))
sheet(wb.create_sheet("Recording list"), "Recording list - each recording once, in the order it is first heard",
      "%d recordings (%d to record). Record each once; the file name is the Audio ID." % (len(uniq), n_tts),
      ["S.No", "Audio ID", "File name", "Hindi script (spoken exactly)", "Plays on", "Duration (s)", "Status"],
      [6, 22, 24, 50, 44, 10, 26], u_rows, dev_cols=(4,), wrap_cols=(4, 5), center=(1, 6))

SFX = [("sfx_play_button", "Play button on the start screen - sounds on the press", "Start"),
       ("sfx_next_button", "Next arrow, and the arrow on the celebration screen (real presses only)", "all pages"),
       ("sfx_pop", "Soft pop as letters, words and pictures appear; the cover title", "cover, 1-5"),
       ("sfx_train_move", "The train's chug as it arrives / leaves", "train pages"),
       ("sfx_whistle", "The train's whistle", "train pages"),
       ("sfx_fb_correct", "Correct answer (also a right word caught in the basket game)", "6-17"),
       ("sfx_fb_incorrect", "Wrong answer (also a wrong word in the basket game)", "6-17"),
       ("sfx_mt_burst", "«मात्रा टोकरी»: a round completed", "17"),
       ("sfx_celebrate", "Celebration screen", "18"),
       ("bgm_lesson", "Background music - from the play button on, ducked under voice and sounds", "all")]
sf = []
for cid, what, pages in SFX:
    path = os.path.join(BUILD, "assets", "Audio", cid + ".ogg")
    if not os.path.exists(path):
        continue
    sf.append(([cid, cid + ".ogg", what, pages, duration(path)], None))
sheet(wb.create_sheet("Sound effects"), "Sound effects and music (supplied - not recorded)",
      "Only the ones the lesson plays.",
      ["Audio ID", "File name", "What it is / where it plays", "Page(s)", "Duration (s)"],
      [20, 22, 80, 12, 11], sf, wrap_cols=(3,), center=(4, 5))

sp = wb.create_sheet("Recording spec")
spec = [("मात्राओं की रेल (ओ / औ) - recording spec", None), ("", None),
        ("Scope", "The whole lesson - start screen, 3 transitions, 18 pages including the «मात्रा टोकरी» game, "
                  "and the celebration."),
        ("Order", "Game order. On each page: what plays as it opens, then per item the 1st / 2nd / 3rd wrong try "
                  "(Hint 1, 2, 3) and the correct answer - recorded from a play-through of the current build. A clip "
                  "heard again later on the same page is listed once, with 'plays again'. Clips that play only for "
                  "other choices of the child are at the end of their page."),
        ("Clips", "%d rows, %d distinct recordings, %d of them to record (see 'Recording list')."
                  % (len(rows), len(uniq), n_tts)),
        ("Status", "'TTS placeholder - record': a machine (Gemini TTS) take - please record it. 'Recorded voice "
                   "(File3)': File3's recorded narrator saying the identical line, copied over - keep it, or record it "
                   "too so the whole lesson is one voice. The three transition lines are the shared kit recordings."),
        ("Folder", "3_BUILD/assets/Audio/<Audio ID>.ogg - the file name is the Audio ID"),
        ("Format", "WAV, 24 kHz mono, 16-bit (it may keep the .ogg name, as the current files do); a short silence "
                   "before and after the voice; about -16 LUFS, peaks at most -1 dBFS"),
        ("Register", "आप form throughout; the celebration line «बहुत बढ़िया, दोस्त! तुमने कमाल कर दिया!» is the standard "
                     "end-screen dialogue and the one allowed तुम form."),
        ("Script", "Record exactly the text in the Hindi script column - it is the build's own text."),
        ("Pauses", "Pages 1-5 show letters, matras, words and pictures ON their words, measured from the recording's "
                   "own pauses - so keep a clear pause at each full stop and comma, e.g. «यह है ओ। इसकी मात्रा देखिए।», "
                   "«घ के साथ ओ की मात्रा लगाने पर, घो बनता है।», «घ। घो। घोड़ा।», «अब खि, लौ और ना को जोड़ने पर, खिलौना बनता है।»."),
        ("After recording", "Put the files in 3_BUILD/assets/Audio/ with the same names and re-run the build "
                            "(PYTHONUTF8=1 python 1_SPEC/build_skill_HI02H11_L02_S04.py): it re-measures the timing "
                            "cues from the new takes and re-stamps the audio so browsers fetch the new files."),
        ("Regenerate", "PYTHONUTF8=1 python 1_SPEC/gen_lesson_vo_manifest.py")]
for a, b_ in spec:
    sp.append([a, b_] if b_ is not None else [a]); r = sp.max_row
    sp.cell(row=r, column=1).font = Font(bold=True, size=14 if r == 1 else 11, color="1F4E79" if b_ is None else "000000")
    if b_ is not None:
        sp.cell(row=r, column=2).font = Font(name=DEV, size=11)
        sp.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        sp.cell(row=r, column=1).alignment = Alignment(vertical="top")
sp.column_dimensions["A"].width = 16; sp.column_dimensions["B"].width = 120

wb.save(OUT)
print("  wrote", OUT)
print("  %d rows in play order, %d distinct recordings (%d to record), %d missing, %d sound effects"
      % (len(rows), len(uniq), n_tts, sum(1 for r in rows if r["status"].startswith("MISSING")), len(sf)))
print("  declared on the card but never played (left out):", [(c, text.get(c, "")) for c in not_played])
